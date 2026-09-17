"""1.1 acceptance: qpdf structure, subset font evidence, original positions."""
import sys,io,json,base64,copy,subprocess,tempfile,logging
from pathlib import Path
root=Path(__file__).resolve().parents[1];sys.path.insert(0,str(root/'native'))
logging.getLogger('pypdf').setLevel(logging.ERROR)
import fitz,pikepdf,worker
from fontTools.ttLib import TTFont
from fontTools import subset
from font_match import CACHE
from font_similarity import recommend,fallback
from story import layout
from qpdf_tools import process
out=root/'tests/output';out.mkdir(exist_ok=True)
checks=[]
def ok(name): checks.append(name);print('PASS',name,flush=True)
# Create a real subset with misleading names. Detection must use contours.
font=TTFont(root/'native/fonts/DejaVuSans.ttf');sub=subset.Subsetter();sub.populate(text='ABCDEFGHIJKLMN');sub.subset(font)
for name in font['name'].names:
 if name.nameID in (1,4,6):name.string='MisleadingSans'.encode('utf-16-be') if name.isUnicode() else b'MisleadingSans'
buf=io.BytesIO();font.save(buf);font.close();blob=buf.getvalue()
import hashlib
key=hashlib.sha256(blob).hexdigest()[:32];CACHE.mkdir(exist_ok=True);(CACHE/(key+'.ttf')).write_bytes(blob);(CACHE/(key+'.json')).write_text(json.dumps({'name':'MisleadingSans','coverage':[ord(c) for c in 'ABCDEFGHIJKLMN']}))
r=recommend(key,'Z','ABCDEFGHIJKLMN');assert r and r[0]['exactGlyphs']>=8 and 'DejaVu' in r[0]['name'],r
assert ord('Z') in fallback(key,'Z')['coverage'];ok('Actual outline matching identifies renamed subset and covers missing Z')
# Multi-line with deliberate 0.6pt letter spacing.
pdf=fitz.open();page=pdf.new_page(width=595,height=842);page.insert_font(fontname='A',fontfile=str(root/'native/fonts/DejaVuSans.ttf'))
for line,text in enumerate(['ABCDEFGHIJKLMN','ABCDEFGHIJKLMN','ABCDEFGHIJKLMN']):
 x=70
 for ch in text:
  page.insert_text((x,100+line*20),ch,fontsize=12,fontname='A');x+=fitz.Font(fontfile=str(root/'native/fonts/DejaVuSans.ttf')).text_length(ch,fontsize=12)+.6
raw=pdf.tobytes();pdf.close();path=out/'v11-fixture.pdf';path.write_bytes(raw)
r=worker.run({'command':'inspect','bytes':base64.b64encode(raw).decode()});(out/'v11-inspect.json').write_text(json.dumps(r))
js="import {pageCandidates} from './src/flow-page-model.mjs';import fs from 'fs';const r=JSON.parse(fs.readFileSync('tests/output/v11-inspect.json'));console.log(JSON.stringify(pageCandidates(r.objects,595,842)[0].model));"
m=json.loads(subprocess.check_output(['node','--input-type=module','-e',js],cwd=root));assert m.get('originalLayout'),m
r=layout(m);assert r['layoutMode']=='原始字位';gs=r['glyphs'];assert all(abs(a['baseline']-b['baseline'])<.001 and abs(a['x']-b['x'])<.001 for a,b in zip(gs,m['originalLayout']['glyphs']))
# Replace B with A and keep all other positions exactly.
m['text']=m['text'][:1]+'A'+m['text'][2:];r=layout(m);assert r['layoutMode']=='原始字位'
assert all(abs(a['baseline']-b['baseline'])<.001 and abs(a['x']-b['x'])<.001 for i,(a,b) in enumerate(zip(r['glyphs'],gs)) if i!=1)
fragment=base64.b64decode(r['fragment']);withpdf=fitz.open(stream=fragment,filetype='pdf');assert 'AACD' in ''.join(withpdf[0].get_text().split());withpdf.close();ok('Original glyph positions survive same-width correction on three spaced lines')
# Style and frame changes must deliberately use full reflow.
changed=copy.deepcopy(m);changed['size']=14;assert layout(changed)['engine']=='MuPDF Story';ok('Font size change exits anchored mode without ignoring formatting')
# qpdf preserves structure and does not change source.
for user in ['secret','']:
 with pikepdf.open(path) as p:
  p.Root['/PageMode']=pikepdf.Name('/UseOutlines');p.docinfo['/Title']='QA document'
  enc=out/('v11-encrypted'+('open' if user else 'restricted')+'.pdf');p.save(enc,encryption=pikepdf.Encryption(owner='owner',user=user,R=6,allow=pikepdf.Permissions(extract=False,modify_other=False)))
 before=enc.read_bytes();target=out/'v11-decrypted.pdf'
 if user:assert process({'command':'qpdf-decrypt','input':str(enc),'output':str(target),'password':'wrong'})['needsPassword']
 r=process({'command':'qpdf-decrypt','input':str(enc),'output':str(target),'password':user});assert r['verified'] and r['pages']==1
 with pikepdf.open(target) as p:assert not p.is_encrypted and p.Root.PageMode=='/UseOutlines' and p.docinfo.Title=='QA document'
 assert enc.read_bytes()==before
ok('qpdf decrypts AES-256 and empty-open-password restrictions; wrong password rejected; source untouched')
(out/'v11-native-report.json').write_text(json.dumps({'checks':checks,'qpdf':pikepdf.__libqpdf_version__},ensure_ascii=False,indent=2))
# Structured exports preserve leading zeros, formulas as plain text, and merges.
from table_export import export
from openpyxl import load_workbook
table={'rows':2,'columns':2,'cells':[{'row':0,'column':0,'colSpan':2,'model':{'text':'00123'}},{'row':1,'column':0,'model':{'text':'=SUM(A1:A2)'}},{'row':1,'column':1,'model':{'text':'中文'}}]}
x=export(table,'xlsx');wb=load_workbook(io.BytesIO(base64.b64decode(x['base64'])));ws=wb.active
assert ws['A1'].value=='00123' and ws['A2'].value=='=SUM(A1:A2)' and ws['A2'].data_type=='s' and str(ws.merged_cells)=='A1:B1'
ok('XLSX exports literal text, leading zeros and merged cells')
# Reopened output: repeated OCR replaces only a recognized Folio layer.
block={'page':1,'text':'OCR SAMPLE','confidence':1,'quad':[[70,700],[170,700],[170,682],[70,682]]}
first=worker.run({'command':'apply','input':str(path),'ocr':[block]});one=base64.b64decode(first['bytes']);second=worker.run({'command':'apply','bytes':base64.b64encode(one).decode(),'ocr':[{**block,'text':'OCR UPDATED'}]});two=base64.b64decode(second['bytes'])
with fitz.open(stream=two,filetype='pdf') as p:
 text=p[0].get_text();assert 'OCR SAMPLE' not in text and text.count('OCR UPDATED')==1 and 'A' in text
ok('Reapplying OCR to reopened Folio output replaces its identified layer once')
(out/'v11-native-report.json').write_text(json.dumps({'checks':checks,'qpdf':pikepdf.__libqpdf_version__},ensure_ascii=False,indent=2))
